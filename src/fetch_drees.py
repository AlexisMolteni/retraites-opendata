"""
Parsers pour les fichiers DREES Panorama des retraites 2025.

Fichiers attendus dans data/raw/external/ :
  drees_vue_ensemble.xlsx     — Vue d'ensemble Graphique 1 (âge conjoncturel H/F/T)
  drees_fiche05_niveau.xlsx   — F05_Tableau 1 (pension brute mensuelle H/F/T 2004-2023)
  drees_fiche07_nouveaux.xlsx — F07_Graphique 1 (pension primo-liquidants vs ensemble)
"""
import re
from pathlib import Path

import openpyxl
import pandas as pd

RAW_EXT = Path("data/raw/external")


def _to_float(val) -> float | None:
    if val is None:
        return None
    s = str(val).replace("nd", "").replace(" ", "").strip()
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def _clean_year(val) -> int | None:
    if val is None:
        return None
    s = re.sub(r"[^0-9]", "", str(val).split(".")[0])[:4]
    return int(s) if len(s) == 4 else None


# ─── Vue d'ensemble Graphique 1 — âge conjoncturel de départ ─────────────────

def parse_age_depart() -> pd.DataFrame:
    """
    Vue d'ensemble_Graphique 1 :
      row 3 (0-idx) : années 2004-2023 aux cols 2-21
      row 4 : Femmes, row 5 : Hommes, row 6 : Ensemble
    """
    wb = openpyxl.load_workbook(RAW_EXT / "drees_vue_ensemble.xlsx", data_only=True, read_only=True)
    ws = wb["Vue d'ensemble_Graphique 1"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    years    = [_clean_year(v) for v in rows[3][2:22]]
    femmes   = [_to_float(v)  for v in rows[4][2:22]]
    hommes   = [_to_float(v)  for v in rows[5][2:22]]
    ensemble = [_to_float(v)  for v in rows[6][2:22]]

    records = []
    for i, yr in enumerate(years):
        if yr is None:
            continue
        for gc, vals in [("F", femmes), ("H", hommes), ("T", ensemble)]:
            v = vals[i] if i < len(vals) else None
            if v is not None:
                records.append({
                    "annee": yr, "regime_id": 99, "genre_code": gc,
                    "indicateur": "AGE_CONJONCTUREL_DEPART",
                    "valeur": v, "unite": "ans", "source_fichier": "DREES_PANORAMA_2025",
                })
    return pd.DataFrame(records)


# ─── Fiche 05 Tableau 1 — niveau de pension brute mensuelle ──────────────────

def parse_niveau_pensions() -> pd.DataFrame:
    """
    F05_Tableau 1 : pension brute mensuelle droit direct y compris majorations.

    Structure (0-indexed openpyxl rows) :
      rows 0-5 : en-têtes multi-niveaux
      rows 6-25 : données 2004-2023
      row 26 : note de bas de page

    Colonnes (0-indexed) :
      1  : année
      2  : DD hors maj brut Ensemble
      3  : DD y.c. maj brut Femmes   ← cible F
      4  : DD y.c. maj brut Hommes   ← cible H
      5  : DD y.c. maj brut Ensemble ← cible T
      6  : DD y.c. maj net  Ensemble
      7  : Pension totale brut Femmes
      8  : Pension totale brut Hommes
      9  : Pension totale brut Ensemble
      10 : Pension totale net  Ensemble
      11+: % évolutions
    """
    wb = openpyxl.load_workbook(RAW_EXT / "drees_fiche05_niveau.xlsx", data_only=True, read_only=True)
    ws = wb["F05_Tableau 1"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    records = []
    for row in rows:
        yr = _clean_year(row[1]) if len(row) > 1 else None
        if yr is None or not (2004 <= yr <= 2023):
            continue
        for gc, ci in [("F", 3), ("H", 4), ("T", 5)]:
            v = _to_float(row[ci]) if len(row) > ci else None
            if v is not None:
                records.append({
                    "annee": yr, "regime_id": 99, "genre_code": gc,
                    "indicateur": "PENSION_BRUTE_MOIS_DD_MAJ",
                    "valeur": v, "unite": "EUR/mois",
                    "source_fichier": "DREES_FICHE05_2025",
                })
    return pd.DataFrame(records)


# ─── Fiche 07 Graphique 1 — pension nouveaux retraités ───────────────────────

def parse_nouveaux_retraites_pension() -> pd.DataFrame:
    """
    F07_Graphique 1 : pension mensuelle brute (€ constants 2023) primo-liquidants et ensemble.

    Structure (0-indexed) :
      En-têtes dans les premières lignes
      Colonnes : 1=année, 2=primo-liquidants, 3=ensemble tous retraités
      Valeurs 'nd' pour 2018-2019 (transition LURA)
    """
    wb = openpyxl.load_workbook(RAW_EXT / "drees_fiche07_nouveaux.xlsx", data_only=True, read_only=True)
    # Le nom de feuille a un espace de fin dans le fichier DREES
    sheet_name = next(n for n in wb.sheetnames if n.strip() == "F07_Graphique 1")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    records = []
    for row in rows:
        yr = _clean_year(row[1]) if len(row) > 1 else None
        if yr is None or not (2004 <= yr <= 2023):
            continue
        for ind, ci in [
            ("PENSION_BRUTE_NOUVEAUX_EUR2023", 2),
            ("PENSION_BRUTE_ENSEMBLE_EUR2023", 3),
        ]:
            v = _to_float(row[ci]) if len(row) > ci else None
            if v is not None:
                records.append({
                    "annee": yr, "regime_id": 99, "genre_code": "T",
                    "indicateur": ind,
                    "valeur": v, "unite": "EUR2023/mois",
                    "source_fichier": "DREES_FICHE07_2025",
                })
    return pd.DataFrame(records)


# ─── Agrégateur ───────────────────────────────────────────────────────────────

def fetch_all_drees(include_existing: bool = False) -> pd.DataFrame:
    """
    Agrège tous les parsers disponibles.

    include_existing=False : uniquement les nouvelles fiches (05 et 07)
    include_existing=True  : ajoute aussi vue_ensemble si disponible
    """
    dfs = [parse_niveau_pensions(), parse_nouveaux_retraites_pension()]
    if include_existing and (RAW_EXT / "drees_vue_ensemble.xlsx").exists():
        dfs.append(parse_age_depart())
    df = pd.concat(dfs, ignore_index=True)
    return df


def fetch_all_drees_initial() -> pd.DataFrame:
    """Chargement initial (première session) — inclut vue_ensemble."""
    dfs: list[pd.DataFrame] = []
    if (RAW_EXT / "drees_vue_ensemble.xlsx").exists():
        dfs.append(parse_age_depart())
    dfs += [parse_niveau_pensions(), parse_nouveaux_retraites_pension()]
    return pd.concat(dfs, ignore_index=True)
